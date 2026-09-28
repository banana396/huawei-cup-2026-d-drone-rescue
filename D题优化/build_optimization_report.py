# AI-assisted code: OpenAI Codex; metadata to be verified by team.
from audit_optimize import *
import openpyxl

def main():
    load=lambda p:json.loads(p.read_text(encoding='utf8'))
    orig=ROOT.parent/'D题求解'/'结果';a=load(orig/'q2_results.json')['selected'];b=load(OUT/'q2_results.json')['selected']
    c=load(orig/'q3_solution.json');d=load(OUT/'q3_solution.json');q4=load(OUT/'q4_results.json');e=Engine()
    qa=Counter(r['model'] for r in a['missions']);qb=Counter(r['model'] for r in b['missions'])
    oldlate=sum(max(0,t-e.boxes[i].expected_s) for r in a['missions'] for i,t in r['delivery_s'].items())
    oldmax=max(max(0,t-e.boxes[i].expected_s) for r in a['missions'] for i,t in r['delivery_s'].items())
    joint=lambda q:max(r['return_s'] for r in q['transport_missions']+q['relay_sorties'])
    pairs=[('Q2架次',len(a['missions']),len(b['missions'])),('Q2完工时间_s',a['metrics']['makespan_s'],b['metrics']['makespan_s']),
           ('Q2能耗_kWh',a['metrics']['energy_kwh'],b['metrics']['energy_kwh']),('Q2最大迟到_s',oldmax,0),('Q2总迟到_s',oldlate,0),
           *[(f'{m}型架次',qa[m],qb[m]) for m in 'ABC'],('Q3运输架次',len(c['transport_missions']),len(d['transport_missions'])),
           ('Q3中继班次',len(c['relay_sorties']),len(d['relay_sorties'])),('Q3联合完工_s',joint(c),joint(d)),
           ('Q3中继能耗_kWh',sum(r['energy_kwh'] for r in c['relay_sorties']),sum(r['energy_kwh'] for r in d['relay_sorties'])),
           ('Q3总能耗_kWh',sum(r['energy_kwh'] for r in c['transport_missions']+c['relay_sorties']),sum(r['energy_kwh'] for r in d['transport_missions']+d['relay_sorties']))]
    csvwrite('before_after_comparison.csv',[{'指标':k,'原方案':x,'新方案':y,'原减新':x-y,'下降比例':(x-y)/x if x else ''} for k,x,y in pairs])
    table='\n'.join(f'| {k} | {x:.6f} | {y:.6f} | {(x-y)/x*100:.2f}% |' for k,x,y in pairs if x)
    use_before=list(csv.DictReader((ROOT/'craft_utilization_before.csv').open(encoding='utf-8-sig')))
    use_after=list(csv.DictReader((ROOT/'craft_utilization_after.csv').open(encoding='utf-8-sig')))
    util='\n'.join(f"| {x['craft']} | {x['missions']} / {y['missions']} | {float(x['utilization']):.1%} / {float(y['utilization']):.1%} | {float(x['last_return_s']):.3f} / {float(y['last_return_s']):.3f} |" for x,y in zip(use_before,use_after))
    q4text=[]
    for k in ('2','3'):
        q=q4['solutions'][k]['balanced_resource_priority'];q4text.append(f"### {k} 组：工作量差异率 {q['work_imbalance_ratio']:.2%}\n")
        for idx,g in enumerate(q['groups'],1):
            r=g['resources'];q4text.append(f"- G{idx}：{', '.join(sorted(g['zones']))}。A/B/C 运输机 {r['A_craft']}/{r['B_craft']}/{r['C_craft']}，电池 {r['A_batteries']}/{r['B_batteries']}/{r['C_batteries']}，中继机/能源组件 {r['relay_craft']}/{r['relay_components']}。")
        q4text.append('\n总需求：'+json.dumps(q['total_required'],ensure_ascii=False)+'。\n\n库存缺口：'+json.dumps({x:v for x,v in q['shortfall'].items() if v},ensure_ascii=False)+'。\n')
    relay=list(csv.DictReader((ROOT/'relay_schedule.csv').open(encoding='utf-8-sig')))
    rt='\n'.join(f"| {r['relay_sortie']} / {r['craft']} | {float(r['takeoff_s']):.3f} | {float(r['lon']):.8f}, {float(r['lat']):.8f}, {float(r['alt_m']):.3f} | {float(r['service_start_s']):.3f}–{float(r['service_end_s']):.3f} | {float(r['return_s']):.3f} | {r['mission_count']} | {float(r['energy_kwh']):.6f} |" for r in relay)
    opt=load(OUT/'q2_results.json')['optimization'];comm=load(OUT/'q3_communication_check.json');cert=load(OUT/'q3_interval_certificate.json')
    margin=load(OUT/'q3_margin_certificate_1db.json');jointcheck=load(OUT/'q3_shared_search_check.json')
    report=f'''# D题算法审计与优化报告

更新日期：2026-09-28。所有新结果位于本目录，原项目 `../D题求解`、原始题目、原 Excel 均保留。公开参考数字未参与求解，未访问或复制网上方案。结果是给定物理模型内、经过约束验收的有限搜索最好方案，不声称全局最优或实飞验证。

## 1. 最终结果与复现基准

Q1 完全保持：18 架次、59.130424 kWh。公共物理代码 `d_common.py`、Q1 求解程序及 Q1 数值文件与原项目 SHA-256 相同。

原 Q2 的 26 架次、12595.021415 s、73.20763745 kWh 已通过原求解程序独立复现，任务记录逐项相同。原 Q3 15 中继架次也逐项复现。24191.316791 s 是中继最后返航，联合完工应为 24325.643849 s。原 Excel 六表 427 行与该版本一致，没有发现 Excel 数字与代码不一致。证据：`baseline_reproduction.json`。

| 指标 | 原方案 | 新方案 | 下降比例 |
|---|---:|---:|---:|
{table}

机型架次增减是资源使用结构变化，不将 A 型增加解释为性能下降。Q2/Q3 的最大迟到、非加权总迟到及加权迟到均为 0。Q3 运输完工 {d['metrics']['makespan_s']:.6f} s，联合完工 {joint(d):.6f} s。

## 2. 最主要的三个瓶颈及证据

**瓶颈一：机型局部能耗选择导致 B 型长队。** 原 `optimize_q2.py` 虽在方案层比较迟到，最终选中的 `selection='energy'` 在每个任务分配时仍首先选择最低能耗机型。因此 14 个 B 型任务均可在物理上由 A 型执行，却集中于两架 B 型。这里的“可执行”指载荷、体积、能耗和从零出发的时限可行，不能直接视为任意插入时刻可行。原 U06 忙碌 12198.251 s、等待 396.771 s，最后返航 12595.021 s，是实际关键链；四架 A 型各只运行一次。详见 `task_model_feasibility.csv`、逐机时间线与利用率表。

**瓶颈二：固定组批把工作集中到两架 C 型。** 原任务组中有 8 个任务只能由 C 型承担，按它们的作业总时长除以两架 C 型，得到固定组批下 makespan 下界 8525.515840 s。这证明仅把 B 任务迁移到 A、保持所有货箱组批不变，无法达到小于 8500 s；这个下界不适用于重新组批后的模型。固定组批搜索得到 26 架次、8951.475141 s、零迟到，再通过货箱重分配、拆分、合并和顺访降至 6374.558828 s。原算法把硬时限货箱与其他货箱分开组批，且仅搜索软任务次序，遗漏了更好的混合组批和全任务排序。

**瓶颈三：Q3 的固定波次与按软任务新派中继。** 原 `solve_q3.py` 明确设置软任务最早准备时刻 6000 s，固定运输实体及电池，逐软任务调用 `launch()`，再等中继就位。它并非完全忽略悬停或多任务共享：第一架中继已经持续覆盖硬任务三波；但软任务没有把需求汇总成共享窗口，造成反复往返及运输串行等待。新方案按所有新版运输轨迹联合搜索站位与时间窗，3 个中继架次分别保障 15、5、2 个运输架次；同一运输架次可以出现在不同中继的保障集合中。

## 3. 电池、时间窗与资源口径审计

核对题目附录2后，原充电模型正确：SOC<0.9 时，充电时间为 T[0.65(0.9−SOC)/0.9+0.35]；SOC≥0.9 时为 T×0.35(1−SOC)/0.1。同机型电池共享，各电池并行充电，不同机型不混用。原代码的飞机和电池是两个资源池，没有永久绑定，也没有重复施加充电桩上限。新代码保持该口径。最终运输库存 A/B/C 为 4/2/2，电池为 6/4/4；两架中继、六组组件，不增加资源。

表中 `start_s` 是准备开始，不是起飞；时间线文件另列起飞=准备开始+固定准备+逐箱装载。医疗期望时刻及首批截止为硬约束，其余期望时刻用于迟到目标，未擅自放宽。所有节点间飞行都仍按逐 DEM 像元最大高程、固定净空、爬升/水平/下降和原能耗计算。

逐机统计的利用率为任务占用时间/全局完工时间，包含准备、装载、飞行和交接。`flight_s` 排除准备和交接。电池等待以该飞机空闲后任一同机型电池首次可用时间核算，属于给定排程下的资源等待诊断；其余空闲记作调度或无剩余任务空闲。题目中所有货箱初始已知，没有外部任务到达时刻，所以不能把该空闲误说成“货物还没到”。

| 无人机 | 原/新架次 | 原/新利用率 | 原/新最后返航(s) |
|---|---:|---:|---:|
{util}

新版关键机变为 U08（C 型），忙碌时间等于 6374.558828 s、无等待；U07 也接近满负荷。这是当前排程的关键链，不是全部可能组批的全局下界。进一步优化须改变 C 型任务组合或分配，而不能仅消除现有关键机空闲。

## 4. 搜索方法、目标顺序与 checkpoints

目标按“全部硬约束可行 → 总迟到 → makespan → 架次 → 能耗”字典序选择。总迟到采用逐箱非加权迟到，同时保留原加权指标。所有最终候选两者均为0，因此不影响零迟到候选之间的排序。

Q2 使用 EDF/期望时限排序、按重量递减装箱、多起点随机化构造、机型迁移、任务交换与重排、逐箱 relocate/swap、merge/split、两区顺访及模拟退火接受机制。接受非改进方案仅用于跳出局部极小；另存的最好方案不会被覆盖丢失。初轮评估 {opt['evaluated']} 个候选、独立验收 {opt['validated']} 个可行候选；重启搜索评估 {opt['refinement']['evaluated']} 个；25/24架次专项评估 {opt['sortie_caps']['evaluated']} 个。种子为 20260928、92871、24928。所有结构可行候选经独立逐段运输检查；容量/时限已失败的候选提前剔除。

阶段记录：固定组批26架次8951.475141 s；重组26架次6425.727504 s；25架次6374.558828 s、72.867419 kWh；24架次同完工、72.890983 kWh。选择24架次意味着为少一次起降增加约0.023565 kWh，符合“架次先于能耗”的规则。完整记录保存在 `checkpoints/` 与搜索日志。

Q3 保持新版Q2的货箱组合、路线及机型，联合重排运输资源与时刻。候选位置由所有节点范围的23×23网格、节点及基地—节点中点生成，离地高度取150/225/300 m的候选，严格不超过300 m；没有硬编码最终坐标。经精确回传和接入筛选，搜索一处持续站位加另一架中继的两段共享服务窗口，随后做任务次序/转场时刻局部改进。搜索可行候选先检查运输、资源和预计算采样覆盖，最终胜出候选再执行独立密采样、连续区间及通信记录验收；中间候选不冒充已通过连续证明的结果。`q3_shared_search_check.json` 记录 {jointcheck['tested']} 次排程尝试及 {jointcheck['transport_feasible']} 个运输可行候选。

## 5. 中继安排

坐标为算法从原始数据独立生成的输出，海拔单位m，时间单位s。组件依次为R-EC-01、02、03。

| 架次/中继机 | 起飞 | 经度、纬度、海拔 | 服务窗口 | 返回 | 保障运输架次个数 | kWh |
|---|---:|---|---|---:|---:|---:|
{rt}

完整服务运输编号、准备时刻、返航SOC在 `relay_schedule.csv`。官方模板没有Q3运输架次独立表，因此其24个运输任务、逐箱交付和资源记录另附 `结果/q3_solution.json`，不可拿Q2时间替代。

## 6. Q4重算结果

先冻结新版Q3的236条通信记录与75个边界备用保障关系，再枚举两组511、三组18660个带标签分配。主表沿用工作量差异率≤20%后的“库存缺口、资源总量、差异率”排序。跨组依赖同一中继架次时，各组需配置独立实体执行相同保障任务，不允许跨组共享。

{chr(10).join(q4text)}
两组缺 A机2、C机1、A电池2、中继机2；三组缺 A机4、C机2、A电池3、C电池1、中继机3。两组三组缺口单位数分别7和13，大于旧均衡方案的2和7。这是冻结更紧凑时序后组间无法共享资源的代价，不是Q2/Q3超库存飞行。若只最小化缺口，存在2/3组缺2/5个单位的极不均衡分区（差异率94.87%/90.93%），已保留在JSON，未在官方主表偷换选择原则。

## 7. 全部硬约束验收

- Q2/Q3各80/80箱；重复0、漏箱0、错区0、载重/体积违规0、SOC违规0、飞机/电池占用冲突0、库存违规0、医疗/首批截止违规0。
- Q2按题意不考察通信；Q3的精确逐格3秒采样无未覆盖点，连续区间证书 {cert['counts']['intervals']} 个初始区间全部认证（含必要细分），名义通信中断区间0。
- 通信记录独立复核 {comm['point_checks']} 个点、{comm['certified_intervals']} 个非边界区间，0项错误；75个≤1e-7 s数值边界均有整段备用路径，总误差带宽 {comm['boundary_total_duration_s']:.9g} s。
- Q4六种配置的冻结关系、分区、组内资源峰值和通信复核通过；超过库存的量明确作为第四问要求的配置缺口报告。
- 交付时刻、错区、重复箱、SOC、飞机重叠、电池重叠六类故障注入全部被检出；通信错误中继和Q4删除依赖亦被检出。
- 原始题目和输入哈希与原项目清单一致；Q1结果及公共物理代码保持原字节。原Excel备份哈希一致。
- 新Excel六表366行逐单元格核对0不符；表名、列名及Q1值保持一致，未新增辅助列。更改视图已渲染检查；修复原Q2/Q4表头异常409.5pt高度，并为Q4长服务区列表换行，未调整任何计算数值。

证据入口：`结果/optimized_validation.json`、`q2_check.json`、`q3_check.json`、`q3_interval_certificate.json`、`q3_communication_check.json`、`q3_boundary_guards.json`、`q4_check.json`、`template_check.json`。检查通过仅适用于各自声明的约束和本次结果哈希。

## 8. 失败候选、风险与未完成的人工作业

首次共享中继筛选采用粗视线采样，遗漏S004遮挡；精确检查拒绝该候选。随后改为逐格视线重新搜索并通过最终验收。失败证据保留于 `原始备份/q3_rejected_coarse_check.json`，未用失败候选填表。

新版满足原题名义通信门限，但额外1 dB测试失败：发现 {margin['counts'].get('point_counterexamples',0)} 个细分检查中的点值反例。旧方案额外1 dB通过，新方案不继承该属性。题给接收门限、衰落裕量、遮挡损耗均未放宽；降低的是旧方案额外添加的鲁棒性要求。名义最小认证裕量约 {cert['min_certified_margin_db']:.8f} dB。没有声称5 m定位/地形扰动可靠，工程使用需另做鲁棒设计。

能耗沿用原模型及其显式假设；未为逼近公开69.13 kWh等参考数字调整常数。当前模型没有全局最优证明，未穷尽任意多站路线或连续空间站位。原17页论文仍包含旧数字，不应与本Excel混用；本次交付是新计算结果及审计报告，论文同步、赛队人工确认和AI工具真实元数据核实尚需完成。本次没有上传或代提交。

## 9. 文件与运行

- `outputs/optimization-20260928/结果提交模板_优化版.xlsx`：新官方六表。
- `before_after_comparison.csv`：全指标原/新对比，差值统一为原减新。
- `craft_timeline_before.csv`、`craft_timeline_after.csv`：逐机完整任务、起飞/返回/下次等待。
- `craft_utilization_before.csv`、`craft_utilization_after.csv`：飞行、空闲、电池等待与利用率。
- `task_model_feasibility.csv`、`box_feasibility.csv`：任务及逐箱重量、体积、时限、距离、高差、可行机型与安全载荷。
- `relay_schedule.csv`：中继资源、位置、共享服务与能耗。
- `结果/q2_results.json`、`q3_solution.json`、`q4_results.json`：正式新方案；Q1使用原样保留的q1_results_v2.json。旧结果引用文件另置 `历史参考结果/`，不作为新版通过证据。

运行 `python run_optimized.py` 可验证当前数值和Excel；`python run_optimized.py --search` 从原输入及旧基线重跑确定性搜索，再完成通信、Q4和数据矩阵。后者会更新本优化目录中的结果，原项目不修改。重新导出Excel使用托管Node执行 `export_optimized.mjs`，然后执行 `check_template.py`。依赖NumPy、Pillow、openpyxl及 @oai/artifact-tool，不需要下载外部求解器。具体运行代码均在本目录。
'''
    (ROOT/'optimization_report.md').write_text(report,encoding='utf8')
    # Compare preserved Excel structure and exact Q1 values/styles.
    old=ROOT/'原始备份/结果提交模板_已填.xlsx';new=ROOT/'outputs/optimization-20260928/结果提交模板_优化版.xlsx'
    wa=openpyxl.load_workbook(old);wb=openpyxl.load_workbook(new)
    preserve={'sheet_names_same':wa.sheetnames==wb.sheetnames,'headers_same':all([c.value for c in wa[s][1]]==[c.value for c in wb[s][1]] for s in wa.sheetnames),
        'q1_cells_same':all(c.value==wb.worksheets[0][c.coordinate].value for row in wa.worksheets[0] for c in row),
        'q1_styles_same':all(c._style==wb.worksheets[0][c.coordinate]._style for row in wa.worksheets[0] for c in row),
        'freeze_panes_same':all(wa[s].freeze_panes==wb[s].freeze_panes for s in wa.sheetnames),
        'merged_ranges_same':all(str(wa[s].merged_cells)==str(wb[s].merged_cells) for s in wa.sheetnames)}
    assert all(preserve.values()),preserve
    write(OUT/'workbook_preservation_check.json',preserve)
    print('report and comparison complete',preserve)
if __name__=='__main__':main()
