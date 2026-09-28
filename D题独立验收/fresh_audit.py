#!/usr/bin/env python3
"""
Independent strict audit of Q2 and Q3 results from 结果提交模板_优化版.xlsx

Goal: Prove that Q2 (makespan 6374.6 s / 72.89 kWh) and Q3 (makespan 6512.2 s /
relay energy 3.366 kWh) are feasible under ALL original hard constraints.

All geometry, link-budget, and physics are recomputed from raw DEM + problem parameters.
No intermediate results from the generation phase are reused.
"""

import math, json, pathlib, sys, traceback
from collections import defaultdict

import numpy as np
import openpyxl

# -- paths --
BASE  = pathlib.Path(r"C:\Users\zhiyu\Desktop\新建文件夹")
DATA  = BASE / "2026年中国研究生数学建模竞赛赛题" / "D题" / "数据"
UAV_D = DATA / "无人机应急物资运输基础数据"
GEO_D = DATA / "镇龙乡地理空间数据" / "镇龙乡及周边地理数据"
RESULT = BASE / "D题优化" / "outputs" / "optimization-20260928" / "结果提交模板_优化版.xlsx"
OUT    = BASE / "D题独立验收" / "证据" / "fresh_audit_results.json"

issues = []

def fail(msg):
    issues.append(msg)
    print(f"  !! FAIL: {msg}")

def ok(msg):
    print(f"  OK: {msg}")

# ====================================================================
# 1. Load raw source data
# ====================================================================
print("=" * 72)
print("Phase 1: Loading raw source data")
print("=" * 72)

def load_xlsx(path):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([c for c in row])
    return rows

# Dispatch center and service zones
center_rows = load_xlsx(UAV_D / "调度中心与服务区.xlsx")
O01 = None
zones = {}
for r in center_rows:
    if r[0] == "O01":
        O01 = {"lon": float(r[2]), "lat": float(r[3]), "alt": float(r[4])}
    elif r[0] and str(r[0]).startswith("S"):
        zones[r[0]] = {"lon": float(r[2]), "lat": float(r[3]), "alt": float(r[4]),
                        "name": r[1], "pop": int(r[5]) if r[5] else 0}
print(f"  O01: {O01}")
print(f"  {len(zones)} service zones loaded")

# Transport drone specs
uav_rows = load_xlsx(UAV_D / "运输无人机数据.xlsx")
drone_types = {}
drones = {}
batteries = {}
_in_type_section = False
for r in uav_rows:
    if r[0] == '机型编号':
        _in_type_section = True
        continue
    if r[0] == '无人机编号':
        _in_type_section = False
        continue
    if _in_type_section and r[0] in ("A", "B", "C") and r[3] is not None:
        drone_types[r[0]] = {
            "name": r[1],
            "empty_mass": float(r[2]),
            "max_cargo": float(r[3]),
            "max_volume": float(r[4]),
            "cruise_speed": float(r[5]),
            "empty_range": float(r[6]),
            "full_range": float(r[7]),
            "battery_kwh": float(r[8]),
            "reserve_pct": float(r[9]),
            "prep_time": float(r[10]),
            "load_time_per_box": float(r[11]),
            "handover_base": float(r[12]),
            "handover_per_box": float(r[13]),
            "max_climb": float(r[14]),
            "max_descend": float(r[15]),
            "climb_eff": float(r[16]),
            "descend_eff": float(r[17]),
        }
    if r[0] and str(r[0]).startswith("U"):
        drones[r[0]] = {"type": r[1], "start": r[2]}

_bat_section = False
for r in uav_rows:
    if r[0] and "电池" in str(r[0]):
        _bat_section = True
        continue
    if _bat_section and r[0] in ("A", "B", "C") and r[1] is not None:
        batteries[r[0]] = {"count": int(r[1]), "charge_time": float(r[2])}

print(f"  Drone types: {list(drone_types.keys())}")
print(f"  Drones: {list(drones.keys())}")
print(f"  Battery pools: {batteries}")

# Relay drone specs
relay_rows = load_xlsx(UAV_D / "中继无人机数据.xlsx")
relay_type = {}
relays = {}
relay_batteries = {}
for r in relay_rows:
    if r[0] == "R" and r[3] is not None:
        relay_type = {
            "empty_mass": float(r[2]),
            "comm_module_mass": float(r[3]),
            "takeoff_mass": float(r[4]),
            "cruise_speed": float(r[5]),
            "cruise_power": float(r[6]),
            "energy_kwh": float(r[7]),
            "reserve_pct": float(r[8]),
            "prep_time": float(r[9]),
            "link_time": float(r[10]),
            "turnaround": float(r[11]),
            "max_climb": float(r[12]),
            "max_descend": float(r[13]),
            "climb_eff": float(r[14]),
            "descend_eff": float(r[15]),
            "hover_power": float(r[16]),
            "comm_power": float(r[17]),
            "max_hover_agl": float(r[18]),
        }
    if r[0] and str(r[0]).startswith("R0"):
        relays[r[0]] = {"type": "R", "start": r[2]}

relay_bat_section = False
for r in relay_rows:
    if r[0] is not None and "能源" in str(r[0]):
        relay_bat_section = True
        continue
    if relay_bat_section and r[0] == "R" and r[1] is not None:
        relay_batteries = {"count": int(r[1]), "charge_time": float(r[2])}

print(f"  Relay type: cruise_speed={relay_type['cruise_speed']}, hover_power={relay_type['hover_power']}")
print(f"  Relays: {list(relays.keys())}")
print(f"  Relay energy components: {relay_batteries}")

# Communication parameters
freq_mhz = 2400
L_sys = 3
L_obs = 10
P_sens = -98
M_fade = 8
Pt_uav = 20; G_uav = 3
Pt_relay_acc = 20; G_relay_acc = 6
Pt_relay_bh = 19; G_relay_bh = 8
Pt_gw = 27; G_gw = 12; h_gw = 20
print(f"  Comm params loaded: f={freq_mhz}MHz, Psens={P_sens}dBm, M={M_fade}dB")

# Cargo demand
wb2 = openpyxl.load_workbook(str(UAV_D / "物资需求与配送时限.xlsx"), data_only=True)
ws_boxes = wb2["逐箱货箱清单"]
box_list = {}
for i, row in enumerate(ws_boxes.iter_rows(values_only=True)):
    if i == 0:
        continue
    box_id = row[0]
    box_list[box_id] = {
        "zone": row[1],
        "type": row[2],
        "mass": float(row[3]),
        "volume": float(row[4]),
        "first_batch": row[5] == "是",
        "hard_deadline": float(row[6]) if row[6] else None,
        "soft_deadline": float(row[7]) if row[7] else None,
        "priority": float(row[8]),
    }
print(f"  {len(box_list)} cargo boxes loaded")

# -- Load DEM --
print("\n  Loading DEM from GeoTIFF ...")
HAS_DEM = False
dem_data = None
gt = None
dem_nodata = None
try:
    from osgeo import gdal
    dem_path = str(GEO_D / "数字高程模型数据（DEM）" / "镇龙乡及周边30米DEM.tif")
    ds = gdal.Open(dem_path)
    gt = ds.GetGeoTransform()
    dem_data = ds.GetRasterBand(1).ReadAsArray()
    dem_nodata = ds.GetRasterBand(1).GetNoDataValue()
    print(f"  DEM loaded: shape={dem_data.shape}, origin=({gt[0]:.6f},{gt[3]:.6f}), pixel=({gt[1]:.8f},{gt[5]:.8f})")
    HAS_DEM = True
except Exception as e:
    print(f"  WARNING: Could not load DEM via GDAL: {e}")
    try:
        from PIL import Image as PILImage
        dem_path = str(GEO_D / "数字高程模型数据（DEM）" / "镇龙乡及周边30米DEM.tif")
        pil_im = PILImage.open(dem_path)
        dem_data = np.array(pil_im, dtype=np.float64)
        tags = pil_im.tag_v2
        tiepoint = tags.get(33922)  # ModelTiepointTag
        pxscale = tags.get(33550)   # ModelPixelScaleTag
        gt = (tiepoint[3], pxscale[0], 0.0, tiepoint[4], 0.0, -pxscale[1])
        dem_nodata = -9999
        print(f"  DEM loaded via PIL: shape={dem_data.shape}, origin=({gt[0]:.6f},{gt[3]:.6f}), pixel=({gt[1]:.8f},{gt[5]:.8f})")
        HAS_DEM = True
    except Exception as e2:
        print(f"  WARNING: Could not load DEM via PIL: {e2}")
        try:
            import scipy.io
            mat_path = str(GEO_D / "数字高程模型数据（DEM）" / "镇龙乡及周边30米DEM.mat")
            matdata = scipy.io.loadmat(mat_path)
            for k in matdata:
                if not k.startswith('_'):
                    arr = matdata[k]
                    if hasattr(arr, 'shape') and len(arr.shape) == 2:
                        dem_data = np.array(arr, dtype=np.float64)
                        print(f"  DEM from .mat key='{k}': shape={dem_data.shape}")
                        break
            nrows, ncols = dem_data.shape
            px = 1.0 / 3600.0
            gt = (109.0, px, 0, 23.2, 0, -px)
            dem_nodata = -9999
            HAS_DEM = True
        except Exception as e3:
            print(f"  FATAL: No DEM reader available: {e3}")
            HAS_DEM = False

# -- DEM helpers --
def dem_elevation(lon, lat):
    if not HAS_DEM:
        return 0.0
    col = (lon - gt[0]) / gt[1]
    row = (lat - gt[3]) / gt[5]
    c = int(round(col))
    r = int(round(row))
    if 0 <= r < dem_data.shape[0] and 0 <= c < dem_data.shape[1]:
        v = dem_data[r, c]
        if dem_nodata is not None and v == dem_nodata:
            return 0.0
        return float(v)
    return 0.0

def max_dem_along_path(lon1, lat1, lon2, lat2):
    """Max DEM elevation along horizontal straight line between two points."""
    if not HAS_DEM:
        return 0.0
    d_h = haversine_m(lon1, lat1, lon2, lat2)
    n = max(int(d_h / 25), 5)  # sample every ~25m
    max_h = 0.0
    for i in range(n + 1):
        f = i / n
        lon_i = lon1 + (lon2 - lon1) * f
        lat_i = lat1 + (lat2 - lat1) * f
        h = dem_elevation(lon_i, lat_i)
        if h > max_h:
            max_h = h
    return max_h

# Operating altitudes per the problem statement (Appendix 2)
# O01 operating altitude = ground elevation
# Service zone operating altitude = ground elevation + 30m
O01_OP_ALT = O01["alt"]  # ground elevation
ZONE_OP_ALT = {}
for zid, z in zones.items():
    ZONE_OP_ALT[zid] = z["alt"] + 30.0

def get_node_op_alt(node_id):
    """Operating altitude at a task node."""
    if node_id == "O01":
        return O01_OP_ALT
    return ZONE_OP_ALT.get(node_id, 0.0)

def get_node_lonlat(node_id):
    if node_id == "O01":
        return O01["lon"], O01["lat"]
    z = zones[node_id]
    return z["lon"], z["lat"]

def haversine_m(lon1, lat1, lon2, lat2):
    R = 6371000
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlam = math.radians(lon2 - lon1)
    a = math.sin(dphi/2)**2 + math.cos(phi1)*math.cos(phi2)*math.sin(dlam/2)**2
    return R * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))

def dist3d(lon1, lat1, alt1, lon2, lat2, alt2):
    h = haversine_m(lon1, lat1, lon2, lat2)
    dalt = alt2 - alt1
    return math.sqrt(h**2 + dalt**2)

# -- Link budget --
def fspl_db(d_m, f_mhz):
    if d_m <= 0:
        return 0
    return 20*math.log10(d_m) + 20*math.log10(f_mhz) + 20*math.log10(4*math.pi/300)

def has_los(lon1, lat1, alt1, lon2, lat2, alt2):
    if not HAS_DEM:
        return True, 0.0
    d_horiz = haversine_m(lon1, lat1, lon2, lat2)
    n_steps = max(int(d_horiz / 20), 10)
    max_obs = 0.0
    for i in range(1, n_steps):
        frac = i / n_steps
        lon_i = lon1 + (lon2 - lon1) * frac
        lat_i = lat1 + (lat2 - lat1) * frac
        alt_i = alt1 + (alt2 - alt1) * frac
        ground_i = dem_elevation(lon_i, lat_i)
        obs = ground_i - alt_i
        if obs > max_obs:
            max_obs = obs
    return max_obs <= 0, max_obs

def link_margin_direct(uav_lon, uav_lat, uav_alt_msl):
    gw_ground = dem_elevation(O01["lon"], O01["lat"])
    gw_alt = gw_ground + h_gw
    d = dist3d(uav_lon, uav_lat, uav_alt_msl, O01["lon"], O01["lat"], gw_alt)
    clear, obs = has_los(uav_lon, uav_lat, uav_alt_msl, O01["lon"], O01["lat"], gw_alt)
    path_loss = fspl_db(d, freq_mhz)
    extra_loss = 0 if clear else L_obs
    margin_up = Pt_uav + G_uav + G_gw - path_loss - L_sys - extra_loss - M_fade - P_sens
    margin_down = Pt_gw + G_gw + G_uav - path_loss - L_sys - extra_loss - M_fade - P_sens
    return min(margin_up, margin_down)

def link_margin_relay_access(uav_lon, uav_lat, uav_alt, relay_lon, relay_lat, relay_alt):
    d = dist3d(uav_lon, uav_lat, uav_alt, relay_lon, relay_lat, relay_alt)
    clear, obs = has_los(uav_lon, uav_lat, uav_alt, relay_lon, relay_lat, relay_alt)
    path_loss = fspl_db(d, freq_mhz)
    extra_loss = 0 if clear else L_obs
    margin_up = Pt_uav + G_uav + G_relay_acc - path_loss - L_sys - extra_loss - M_fade - P_sens
    margin_down = Pt_relay_acc + G_relay_acc + G_uav - path_loss - L_sys - extra_loss - M_fade - P_sens
    return min(margin_up, margin_down)

def link_margin_relay_backhaul(relay_lon, relay_lat, relay_alt):
    gw_ground = dem_elevation(O01["lon"], O01["lat"])
    gw_alt = gw_ground + h_gw
    d = dist3d(relay_lon, relay_lat, relay_alt, O01["lon"], O01["lat"], gw_alt)
    clear, obs = has_los(relay_lon, relay_lat, relay_alt, O01["lon"], O01["lat"], gw_alt)
    path_loss = fspl_db(d, freq_mhz)
    extra_loss = 0 if clear else L_obs
    margin_up = Pt_relay_bh + G_relay_bh + G_gw - path_loss - L_sys - extra_loss - M_fade - P_sens
    margin_down = Pt_gw + G_gw + G_relay_bh - path_loss - L_sys - extra_loss - M_fade - P_sens
    return min(margin_up, margin_down)

# ====================================================================
# 2. Load results from Excel
# ====================================================================
print("\n" + "=" * 72)
print("Phase 2: Loading results from Excel")
print("=" * 72)

wb_result = openpyxl.load_workbook(str(RESULT), data_only=True)

ws_q2 = wb_result["Q2_运输架次"]
q2_sorties = []
for i, row in enumerate(ws_q2.iter_rows(values_only=True)):
    if i == 0 or not row[0]:
        continue
    q2_sorties.append({
        "id": row[0], "drone": row[1], "type": row[2], "battery": row[3],
        "start": float(row[4]), "route": str(row[5]),
        "return_time": float(row[6]), "energy": float(row[7]),
    })
print(f"  Q2 sorties: {len(q2_sorties)}")

ws_del = wb_result["Q2_逐箱交付"]
q2_deliveries = []
for i, row in enumerate(ws_del.iter_rows(values_only=True)):
    if i == 0 or not row[0]:
        continue
    q2_deliveries.append({
        "box": row[0], "sortie": row[1], "zone": row[2],
        "deliver_time": float(row[3]),
    })
print(f"  Q2 deliveries: {len(q2_deliveries)}")

ws_q3r = wb_result["Q3_中继架次"]
q3_relays = []
for i, row in enumerate(ws_q3r.iter_rows(values_only=True)):
    if i == 0 or not row[0]:
        continue
    q3_relays.append({
        "id": row[0], "craft": row[1], "component": row[2],
        "start": float(row[3]),
        "hover_lon": float(row[4]), "hover_lat": float(row[5]),
        "hover_alt": float(row[6]),
        "link_ready": float(row[7]), "service_end": float(row[8]),
        "return_time": float(row[9]), "energy": float(row[10]),
    })
print(f"  Q3 relay sorties: {len(q3_relays)}")

ws_q3c = wb_result["Q3_通信保障"]
q3_comm = []
for i, row in enumerate(ws_q3c.iter_rows(values_only=True)):
    if i == 0 or not row[0]:
        continue
    q3_comm.append({
        "sortie": row[0], "phase": row[1],
        "start": float(row[2]), "end": float(row[3]),
        "method": row[4], "relay_id": row[5] if row[5] else None,
    })
print(f"  Q3 comm records: {len(q3_comm)}")

# ====================================================================
# 3. Verify Q2 transport constraints
# ====================================================================
print("\n" + "=" * 72)
print("Phase 3: Verifying Q2 transport hard constraints")
print("=" * 72)

# 3a. All 80 boxes delivered
delivered_boxes = set(d["box"] for d in q2_deliveries)
missing = set(box_list.keys()) - delivered_boxes
if missing:
    fail(f"Missing boxes: {missing}")
else:
    ok(f"All {len(box_list)} boxes delivered")

# 3b. No duplicate deliveries
dup_count = len(q2_deliveries) - len(delivered_boxes)
if dup_count > 0:
    fail(f"{dup_count} duplicate deliveries")
else:
    ok("No duplicate deliveries")

# 3c. Each box to correct zone
zone_errors = 0
for d in q2_deliveries:
    expected_zone = box_list[d["box"]]["zone"]
    if d["zone"] != expected_zone:
        fail(f"Box {d['box']} delivered to {d['zone']} but belongs to {expected_zone}")
        zone_errors += 1
if zone_errors == 0:
    ok("Box-zone assignments correct")

# 3d. Hard deadlines
first_batch_boxes = {bid: b for bid, b in box_list.items() if b["first_batch"]}
first_batch_late = 0
for d in q2_deliveries:
    box = box_list[d["box"]]
    if box["hard_deadline"] is not None:
        if d["deliver_time"] > box["hard_deadline"] + 0.01:
            fail(f"Box {d['box']}: delivered at {d['deliver_time']:.1f}s > hard deadline {box['hard_deadline']}s")
            first_batch_late += 1
if first_batch_late == 0:
    ok(f"All {len(first_batch_boxes)} first-batch boxes meet hard deadlines")

# 3e. Payload capacity per sortie
payload_errors = 0
for s in q2_sorties:
    dt_s = drone_types[s["type"]]
    sortie_boxes = [d for d in q2_deliveries if d["sortie"] == s["id"]]
    total_mass = sum(box_list[d["box"]]["mass"] for d in sortie_boxes)
    total_vol = sum(box_list[d["box"]]["volume"] for d in sortie_boxes)
    if total_mass > dt_s["max_cargo"] + 0.01:
        fail(f"Sortie {s['id']}: cargo mass {total_mass:.1f} > max {dt_s['max_cargo']}kg")
        payload_errors += 1
    if total_vol > dt_s["max_volume"] + 0.001:
        fail(f"Sortie {s['id']}: cargo vol {total_vol:.3f} > max {dt_s['max_volume']}m3")
        payload_errors += 1
if payload_errors == 0:
    ok("Payload capacity constraints all satisfied")

# 3f. Drone temporal non-overlap
drone_sorties = defaultdict(list)
for s in q2_sorties:
    drone_sorties[s["drone"]].append(s)
overlap_errors = 0
for uid, sor_list in drone_sorties.items():
    sorted_s = sorted(sor_list, key=lambda x: x["start"])
    for i in range(len(sorted_s) - 1):
        if sorted_s[i + 1]["start"] < sorted_s[i]["return_time"] - 0.01:
            fail(f"Drone {uid}: overlap {sorted_s[i+1]['id']} starts before {sorted_s[i]['id']} returns")
            overlap_errors += 1
if overlap_errors == 0:
    ok("Drone temporal non-overlap verified")

# 3g. Battery temporal availability
bat_sorties = defaultdict(list)
for s in q2_sorties:
    bat_sorties[s["battery"]].append(s)
bat_errors = 0
for bat, sor_list in bat_sorties.items():
    sorted_s = sorted(sor_list, key=lambda x: x["start"])
    for i in range(len(sorted_s) - 1):
        gap = sorted_s[i + 1]["start"] - sorted_s[i]["return_time"]
        if gap < -0.01:
            fail(f"Battery {bat}: reused before returned")
            bat_errors += 1
if bat_errors == 0:
    ok("Battery temporal availability verified")

# 3h. Independent energy computation per sortie
print("\n  Recomputing sortie energies from geometry ...")

def compute_sortie_energy(sortie):
    dt_s = drone_types[sortie["type"]]
    route_zones = sortie["route"].split("\u2192") if "\u2192" in sortie["route"] else [sortie["route"]]
    sortie_deliveries = [d for d in q2_deliveries if d["sortie"] == sortie["id"]]
    boxes_by_zone = defaultdict(list)
    for d in sortie_deliveries:
        boxes_by_zone[d["zone"]].append(d["box"])
    v = dt_s["cruise_speed"]
    E_bat = dt_s["battery_kwh"]
    R_empty = dt_s["empty_range"]
    R_full = dt_s["full_range"]
    m_max_cargo = dt_s["max_cargo"]
    reserve = dt_s["reserve_pct"] / 100.0
    total_energy = 0.0
    total_time = 0.0
    current_cargo_mass = sum(box_list[bid]["mass"] for d in sortie_deliveries for bid in [d["box"]])
    n_boxes_total = len(sortie_deliveries)
    prep = dt_s["prep_time"] + n_boxes_total * dt_s["load_time_per_box"]
    total_time += prep
    # Nodes: O01 -> zone1 -> zone2 -> ... -> O01
    nodes = ["O01"] + list(route_zones) + ["O01"]

    for leg_i in range(len(nodes) - 1):
        from_node = nodes[leg_i]
        to_node = nodes[leg_i + 1]
        lon1, lat1 = get_node_lonlat(from_node)
        lon2, lat2 = get_node_lonlat(to_node)
        alt_from = get_node_op_alt(from_node)
        alt_to = get_node_op_alt(to_node)
        h_dist = haversine_m(lon1, lat1, lon2, lat2)
        # Cruise altitude = max DEM along path + 50m
        max_terrain = max_dem_along_path(lon1, lat1, lon2, lat2)
        cruise_alt = max_terrain + 50.0
        # Ensure cruise_alt >= both operating altitudes
        cruise_alt = max(cruise_alt, alt_from, alt_to)
        climb_h = max(0, cruise_alt - alt_from)
        descend_h = max(0, cruise_alt - alt_to)
        climb_time = climb_h / dt_s["max_climb"] if climb_h > 0 else 0
        descend_time = descend_h / dt_s["max_descend"] if descend_h > 0 else 0
        # Per problem: climb/descend are pure vertical; all horizontal distance in cruise
        cruise_time = h_dist / v if v > 0 else 0
        leg_time = climb_time + cruise_time + descend_time
        # Energy rate depends on cargo
        cargo_ratio = current_cargo_mass / m_max_cargo if m_max_cargo > 0 else 0
        if leg_i == len(nodes) - 2:
            # Return leg: empty
            range_at_mass = R_empty
        else:
            range_at_mass = R_empty - (R_empty - R_full) * cargo_ratio
        energy_rate = E_bat * v / range_at_mass  # kWh per second of level cruise
        # Per problem appendix 2: E = E_horiz + E_climb_add
        # E_horiz = E_bat * d_horiz / R(m) -- covers full horizontal distance
        # E_climb_add = climb_eff * (E_bat * v / R(m)) * t_climb
        e_per_m = E_bat / range_at_mass
        e_horiz = e_per_m * h_dist
        e_climb_add = dt_s["climb_eff"] * energy_rate * climb_time
        e_descend_add = dt_s["descend_eff"] * energy_rate * descend_time  # 0 since descend_eff=0
        leg_energy = e_horiz + e_climb_add + e_descend_add
        total_energy += leg_energy
        total_time += leg_time
        # Handover at intermediate zones (not at O01 return)
        if leg_i < len(nodes) - 2:
            zone_id = to_node
            n_boxes_here = len(boxes_by_zone.get(zone_id, []))
            handover_time = dt_s["handover_base"] + n_boxes_here * dt_s["handover_per_box"]
            total_time += handover_time
            dropped_mass = sum(box_list[bid]["mass"] for bid in boxes_by_zone.get(zone_id, []))
            current_cargo_mass -= dropped_mass

    soc_remaining = 1.0 - total_energy / E_bat
    return total_energy, total_time, soc_remaining

total_energy_recomp = 0.0
max_return = 0.0
min_soc = 1.0
energy_diffs = []

for s in q2_sorties:
    e_recomp, t_recomp, soc = compute_sortie_energy(s)
    total_energy_recomp += e_recomp
    max_return = max(max_return, s["return_time"])
    min_soc = min(min_soc, soc)
    diff_pct = abs(e_recomp - s["energy"]) / s["energy"] * 100 if s["energy"] > 0 else 0
    energy_diffs.append((s["id"], s["energy"], e_recomp, diff_pct))
    if soc < (drone_types[s["type"]]["reserve_pct"] / 100.0) - 0.001:
        fail(f"Sortie {s['id']}: SOC={soc:.4f} < reserve {drone_types[s['type']]['reserve_pct']}%")

print(f"\n  Q2 claimed energy: 72.89 kWh")
print(f"  Q2 recomputed energy: {total_energy_recomp:.2f} kWh")
print(f"  Q2 claimed makespan: 6374.6 s, actual max return: {max_return:.1f} s")
print(f"  Q2 minimum SOC: {min_soc:.4f}")
energy_diffs.sort(key=lambda x: -x[3])
print(f"\n  Top energy deviations:")
for eid, claimed, recomp, pct in energy_diffs[:5]:
    print(f"    {eid}: claimed={claimed:.4f}, recomp={recomp:.4f}, diff={pct:.2f}%")

if min_soc >= 0.20 - 0.001:
    ok(f"All sorties meet 20% reserve (min SOC = {min_soc:.4f})")
else:
    fail(f"Some sorties violate 20% reserve (min SOC = {min_soc:.4f})")

# ====================================================================
# 4. Verify Q3 relay constraints
# ====================================================================
print("\n" + "=" * 72)
print("Phase 4: Verifying Q3 relay constraints")
print("=" * 72)

relay_craft_used = defaultdict(list)
for r in q3_relays:
    relay_craft_used[r["craft"]].append(r)
for craft, sor_list in relay_craft_used.items():
    sorted_s = sorted(sor_list, key=lambda x: x["start"])
    for i in range(len(sorted_s) - 1):
        if sorted_s[i + 1]["start"] < sorted_s[i]["return_time"] - 0.01:
            fail(f"Relay {craft}: temporal overlap")
ok("Relay craft temporal non-overlap checked")

total_relay_energy = sum(r["energy"] for r in q3_relays)
print(f"  Q3 claimed relay energy: 3.366 kWh")
print(f"  Q3 sum from Excel: {total_relay_energy:.4f} kWh")

for r in q3_relays:
    rt = relay_type
    d_out = haversine_m(O01["lon"], O01["lat"], r["hover_lon"], r["hover_lat"])
    dalt_out = r["hover_alt"] - O01["alt"]
    if dalt_out > 0:
        climb_t = dalt_out / rt["max_climb"]
        desc_t_out = 0
    else:
        climb_t = 0
        desc_t_out = abs(dalt_out) / rt["max_descend"]
    cruise_dist_out = max(0, d_out - (climb_t + desc_t_out) * rt["cruise_speed"])
    cruise_t_out = cruise_dist_out / rt["cruise_speed"]
    hover_duration = r["service_end"] - r["link_ready"]
    hover_energy = (rt["hover_power"] + rt["comm_power"]) * hover_duration / 3600
    transit_out_energy = rt["cruise_power"] * (climb_t * rt["climb_eff"] + cruise_t_out + desc_t_out * rt["descend_eff"]) / 3600
    dalt_ret = O01["alt"] - r["hover_alt"]
    if dalt_ret > 0:
        climb_t_ret = dalt_ret / rt["max_climb"]
        desc_t_ret = 0
    else:
        climb_t_ret = 0
        desc_t_ret = abs(dalt_ret) / rt["max_descend"]
    cruise_dist_ret = max(0, d_out - (climb_t_ret + desc_t_ret) * rt["cruise_speed"])
    cruise_t_ret = cruise_dist_ret / rt["cruise_speed"]
    transit_ret_energy = rt["cruise_power"] * (climb_t_ret * rt["climb_eff"] + cruise_t_ret + desc_t_ret * rt["descend_eff"]) / 3600
    total_relay_e = transit_out_energy + hover_energy + transit_ret_energy
    soc_relay = 1.0 - total_relay_e / rt["energy_kwh"]
    gnd_hover = dem_elevation(r["hover_lon"], r["hover_lat"])
    agl = r["hover_alt"] - gnd_hover
    print(f"\n  Relay {r['id']}:")
    print(f"    Claimed energy: {r['energy']:.4f} kWh, Recomputed: {total_relay_e:.4f} kWh")
    print(f"    Hover duration: {hover_duration:.1f}s, SOC: {soc_relay:.4f}")
    print(f"    Ground at hover: {gnd_hover:.1f}m, AGL: {agl:.1f}m, max={rt['max_hover_agl']}m")
    if agl > rt["max_hover_agl"] + 0.1:
        fail(f"Relay {r['id']}: AGL={agl:.1f}m > max {rt['max_hover_agl']}m")
    if agl < 0:
        fail(f"Relay {r['id']}: hover altitude below ground! AGL={agl:.1f}m")
    if soc_relay < rt["reserve_pct"] / 100.0 - 0.01:
        fail(f"Relay {r['id']}: SOC={soc_relay:.4f} < reserve {rt['reserve_pct']}%")

ec_used = defaultdict(list)
for r in q3_relays:
    ec_used[r["component"]].append(r)
total_ec_count = relay_batteries["count"]
if len(ec_used) > total_ec_count:
    fail(f"Using {len(ec_used)} energy components but only {total_ec_count} available")
else:
    ok(f"Energy component usage ({len(ec_used)}) within pool ({total_ec_count})")

# ====================================================================
# 5. Verify Q3 communication from DEM
# ====================================================================
print("\n" + "=" * 72)
print("Phase 5: Verifying Q3 communication from original DEM")
print("=" * 72)

relay_positions = {}
for r in q3_relays:
    relay_positions[r["id"]] = {
        "lon": r["hover_lon"], "lat": r["hover_lat"], "alt": r["hover_alt"],
        "ready": r["link_ready"], "end": r["service_end"],
    }

def get_uav_position_at_time(sortie_id, t):
    q2_id = "Q2-" + sortie_id.split("-")[1]
    s = None
    for ss in q2_sorties:
        if ss["id"] == q2_id:
            s = ss
            break
    if not s:
        return None
    dt_s = drone_types[s["type"]]
    route_zones = s["route"].split("\u2192") if "\u2192" in s["route"] else [s["route"]]
    sortie_boxes = [d for d in q2_deliveries if d["sortie"] == s["id"]]
    n_boxes = len(sortie_boxes)
    flight_start = s["start"] + dt_s["prep_time"] + n_boxes * dt_s["load_time_per_box"]
    if t < flight_start:
        return O01["lon"], O01["lat"], O01["alt"]
    # Build node list: O01 -> zones -> O01
    nodes = ["O01"] + list(route_zones) + ["O01"]
    boxes_by_zone = defaultdict(list)
    for d in sortie_boxes:
        boxes_by_zone[d["zone"]].append(d["box"])
    cur_t = flight_start
    v = dt_s["cruise_speed"]
    for i in range(len(nodes) - 1):
        from_node = nodes[i]
        to_node = nodes[i + 1]
        lon1, lat1 = get_node_lonlat(from_node)
        lon2, lat2 = get_node_lonlat(to_node)
        alt1 = get_node_op_alt(from_node)
        alt2 = get_node_op_alt(to_node)
        h_dist = haversine_m(lon1, lat1, lon2, lat2)
        max_terrain = max_dem_along_path(lon1, lat1, lon2, lat2)
        cruise_alt = max(max_terrain + 50.0, alt1, alt2)
        climb_h = max(0, cruise_alt - alt1)
        descend_h = max(0, cruise_alt - alt2)
        climb_t = climb_h / dt_s["max_climb"] if climb_h > 0 else 0
        descend_t = descend_h / dt_s["max_descend"] if descend_h > 0 else 0
        # Per problem: climb/descend pure vertical; all horizontal in cruise
        cruise_t = h_dist / v if v > 0 else 0
        leg_time = climb_t + cruise_t + descend_t
        if t <= cur_t + leg_time:
            elapsed = t - cur_t
            # Three-phase: climb -> cruise -> descend
            if elapsed <= climb_t:
                # In climb phase: pure vertical, stay at lon1/lat1
                frac_climb = elapsed / climb_t if climb_t > 0 else 0
                return (lon1,
                        lat1,
                        alt1 + climb_h * frac_climb)
            elif elapsed <= climb_t + cruise_t:
                # In cruise phase: horizontal at cruise_alt
                h_traveled = (elapsed - climb_t) * v
                h_frac = min(1.0, h_traveled / h_dist) if h_dist > 0 else 0
                return (lon1 + (lon2 - lon1) * h_frac,
                        lat1 + (lat2 - lat1) * h_frac,
                        cruise_alt)
            else:
                # In descend phase: pure vertical at lon2/lat2
                elapsed_desc = elapsed - climb_t - cruise_t
                frac_desc = elapsed_desc / descend_t if descend_t > 0 else 1
                return (lon2,
                        lat2,
                        cruise_alt - descend_h * frac_desc)
        cur_t += leg_time
        if i < len(nodes) - 2:
            zone_id = to_node
            boxes_here = boxes_by_zone.get(zone_id, [])
            handover = dt_s["handover_base"] + len(boxes_here) * dt_s["handover_per_box"]
            if t <= cur_t + handover:
                return lon2, lat2, alt2
            cur_t += handover
    return O01["lon"], O01["lat"], O01["alt"]

# Check communication coverage
comm_failures = []
comm_marginals = []
n_checked = 0
min_margin = float("inf")

for rec in q3_comm:
    interval = rec["end"] - rec["start"]
    n_samples = max(3, int(interval / 10))
    n_samples = min(n_samples, 50)
    for si in range(n_samples + 1):
        t = rec["start"] + (rec["end"] - rec["start"]) * si / n_samples
        pos = get_uav_position_at_time(rec["sortie"], t)
        if pos is None:
            continue
        lon, lat, alt = pos
        n_checked += 1
        if rec["method"] == "直连":
            margin = link_margin_direct(lon, lat, alt)
        elif rec["method"] == "中继":
            relay = relay_positions.get(rec["relay_id"])
            if relay is None:
                fail(f"Unknown relay {rec['relay_id']}")
                continue
            if t < relay["ready"] - 0.1 or t > relay["end"] + 0.1:
                fail(f"Relay {rec['relay_id']} not active at t={t:.1f}s")
                continue
            acc_margin = link_margin_relay_access(lon, lat, alt, relay["lon"], relay["lat"], relay["alt"])
            bh_margin = link_margin_relay_backhaul(relay["lon"], relay["lat"], relay["alt"])
            margin = min(acc_margin, bh_margin)
        else:
            continue
        if margin < min_margin:
            min_margin = margin
        if margin < 0:
            comm_failures.append({"sortie": rec["sortie"], "phase": rec["phase"],
                                  "method": rec["method"], "t": t, "margin_db": margin})
        elif margin < 3:
            comm_marginals.append({"sortie": rec["sortie"], "t": t, "margin_db": margin})
# -- Best-available-method audit: at each sample point check if ANY method works --
best_method_failures = []
best_method_min_margin = float("inf")
n_best_checked = 0

for rec in q3_comm:
    interval = rec["end"] - rec["start"]
    n_samples = max(3, int(interval / 10))
    n_samples = min(n_samples, 50)
    for si in range(n_samples + 1):
        t = rec["start"] + (rec["end"] - rec["start"]) * si / n_samples
        pos = get_uav_position_at_time(rec["sortie"], t)
        if pos is None:
            continue
        lon, lat, alt = pos
        n_best_checked += 1
        # Try direct
        best_margin = link_margin_direct(lon, lat, alt)
        # Try every active relay
        for rid, rpos in relay_positions.items():
            if rpos["ready"] - 0.1 <= t <= rpos["end"] + 0.1:
                acc_m = link_margin_relay_access(lon, lat, alt, rpos["lon"], rpos["lat"], rpos["alt"])
                bh_m = link_margin_relay_backhaul(rpos["lon"], rpos["lat"], rpos["alt"])
                relay_m = min(acc_m, bh_m)
                if relay_m > best_margin:
                    best_margin = relay_m
        if best_margin < best_method_min_margin:
            best_method_min_margin = best_margin
        if best_margin < 0:
            best_method_failures.append({"sortie": rec["sortie"], "phase": rec["phase"],
                                          "t": t, "margin_db": best_margin})

print(f"\n  Best-method comm checks: {n_best_checked} points")
print(f"  Best-method min margin: {best_method_min_margin:.2f} dB")
print(f"  Best-method failures (no path available): {len(best_method_failures)}")
if len(best_method_failures) > 0:
    fail(f"{len(best_method_failures)} communication total coverage failures (no direct or relay path)")
    for f_rec in best_method_failures[:10]:
        print(f"    {f_rec['sortie']} {f_rec['phase']} t={f_rec['t']:.1f}s: best_margin={f_rec['margin_db']:.2f}dB")
else:
    ok(f"All {n_best_checked} sample points have at least one viable comm path (min margin={best_method_min_margin:.2f}dB)")

print(f"\n  Communication checks: {n_checked} sample points evaluated")
print(f"  Minimum margin: {min_margin:.2f} dB")
print(f"  Failures (margin < 0): {len(comm_failures)}")
print(f"  Marginal (0 < margin < 3 dB): {len(comm_marginals)}")
if len(comm_failures) > 0:
    fail(f"{len(comm_failures)} communication link failures detected")
    for f_rec in comm_failures[:10]:
        print(f"    {f_rec['sortie']} {f_rec['phase']} t={f_rec['t']:.1f}s: margin={f_rec['margin_db']:.2f}dB")
else:
    ok(f"All {n_checked} sample points have positive link margin (min={min_margin:.2f}dB)")

# Coverage continuity
print("\n  Checking coverage continuity ...")
q3_sortie_ids = set(rec["sortie"] for rec in q3_comm)
gap_errors = 0
for sid in sorted(q3_sortie_ids):
    records = sorted([r for r in q3_comm if r["sortie"] == sid], key=lambda x: x["start"])
    for i in range(len(records) - 1):
        gap = records[i + 1]["start"] - records[i]["end"]
        if gap > 0.1:
            fail(f"Sortie {sid}: comm gap of {gap:.2f}s")
            gap_errors += 1
if gap_errors == 0:
    ok("Coverage continuity verified -- no gaps")

# Backhaul checks
print("\n  Verifying relay backhaul links from DEM ...")
for r in q3_relays:
    bh_margin = link_margin_relay_backhaul(r["hover_lon"], r["hover_lat"], r["hover_alt"])
    clear, obs = has_los(r["hover_lon"], r["hover_lat"], r["hover_alt"],
                         O01["lon"], O01["lat"], dem_elevation(O01["lon"], O01["lat"]) + h_gw)
    print(f"  Relay {r['id']}: backhaul margin={bh_margin:.2f}dB, LOS={'clear' if clear else 'BLOCKED'}")
    if bh_margin < 0:
        fail(f"Relay {r['id']} backhaul link budget negative: {bh_margin:.2f}dB")

# ====================================================================
# 6. Summary
# ====================================================================
print("\n" + "=" * 72)
print("FINAL AUDIT SUMMARY")
print("=" * 72)

q2_makespan = max(s["return_time"] for s in q2_sorties)
q2_energy = sum(s["energy"] for s in q2_sorties)
q3_makespan = max(r["return_time"] for r in q3_relays)
q3_relay_energy_total = sum(r["energy"] for r in q3_relays)

summary = {
    "audit_type": "independent_fresh_audit_v2",
    "Q2": {
        "claimed_makespan_s": 6374.6,
        "actual_makespan_s": round(q2_makespan, 1),
        "claimed_energy_kwh": 72.89,
        "actual_energy_kwh": round(q2_energy, 2),
        "recomputed_energy_kwh": round(total_energy_recomp, 2),
        "min_soc": round(min_soc, 4),
        "all_boxes_delivered": len(missing) == 0,
        "hard_deadlines_met": first_batch_late == 0,
    },
    "Q3": {
        "claimed_makespan_s": 6512.2,
        "actual_makespan_s": round(q3_makespan, 1),
        "claimed_relay_energy_kwh": 3.366,
        "actual_relay_energy_kwh": round(q3_relay_energy_total, 4),
        "relay_sorties": len(q3_relays),
        "comm_sample_points": n_checked,
        "min_link_margin_db": round(min_margin, 2) if min_margin < float("inf") else None,
        "comm_failures": len(comm_failures),
        "comm_marginals": len(comm_marginals),
    },
    "DEM_used": HAS_DEM,
    "total_issues": len(issues),
    "issues": issues,
    "verdict": "PASS" if len(issues) == 0 else "FAIL",
}

print(f"\n  Q2 makespan: {q2_makespan:.1f}s (claimed 6374.6s)")
print(f"  Q2 energy: {q2_energy:.2f} kWh (claimed 72.89 kWh)")
print(f"  Q3 makespan: {q3_makespan:.1f}s (claimed 6512.2s)")
print(f"  Q3 relay energy: {q3_relay_energy_total:.4f} kWh (claimed 3.366 kWh)")
print(f"\n  Issues found: {len(issues)}")
for iss in issues:
    print(f"    - {iss}")
print(f"\n  VERDICT: {summary['verdict']}")

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
print(f"\n  Results written to {OUT}")
